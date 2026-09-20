import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize 16:9 widescreen presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_layout)

# Colors
BG_COLOR       = RGBColor(241, 248, 253)  # Soft Ice Blue background #F1F8FD
WHITE          = RGBColor(255, 255, 255)
NAVY_DEEP      = RGBColor(12, 35, 64)     # #0C2340
NAVY_TEXT      = RGBColor(16, 42, 67)     # #102A43
SLATE_TEXT     = RGBColor(71, 85, 105)    # #475569
POLAR_BLUE     = RGBColor(2, 132, 199)    # #0284C7
SKY_TINT       = RGBColor(224, 242, 254)  # #E0F2FE
BORDER_BLUE    = RGBColor(186, 230, 253)  # #BAE6FD
BORDER_CARD    = RGBColor(203, 222, 235)
RED_HEADER     = RGBColor(220, 38, 38)    # #DC2626
RED_TINT       = RGBColor(254, 242, 242)
ORANGE_CALLOUT = RGBColor(254, 243, 199)  # #FEF3C7
ORANGE_BORDER  = RGBColor(245, 158, 11)   # #F59E0B
DARK_HEADER    = RGBColor(15, 39, 68)     # #0F2744
GREEN_ACCENT   = RGBColor(5, 150, 105)

FONT_FAMILY    = "Segoe UI"

# Background
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = BG_COLOR
bg.line.fill.background()

# ---------------------------------------------------------
# 1. TOP HEADER
# ---------------------------------------------------------
# Top Left CompileX Pill
cx_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.25), Inches(1.5), Inches(0.42))
cx_pill.fill.solid()
cx_pill.fill.fore_color.rgb = WHITE
cx_pill.line.color.rgb = BORDER_CARD
cx_pill.line.width = Pt(1.2)
tf = cx_pill.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.text = "CompileX"
p.font.size = Pt(13)
p.font.bold = True
p.font.name = FONT_FAMILY
p.font.color.rgb = NAVY_TEXT
p.alignment = PP_ALIGN.CENTER

# Top Right SIH Logo
sih_path = os.path.abspath("assets/presentation/sih_logo.png")
if os.path.exists(sih_path):
    slide.shapes.add_picture(sih_path, Inches(11.45), Inches(0.18), width=Inches(1.2), height=Inches(0.56))

# Title & Subtitle in Center
t_box = slide.shapes.add_textbox(Inches(2.5), Inches(0.15), Inches(8.333), Inches(0.70))
tf = t_box.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p1 = tf.paragraphs[0]
p1.text = "POLAR-AI"
p1.font.size = Pt(24)
p1.font.bold = True
p1.font.name = FONT_FAMILY
p1.font.color.rgb = NAVY_DEEP
p1.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = "Mission Continuity Intelligence for Extreme Environments"
p2.font.size = Pt(11)
p2.font.bold = True
p2.font.name = FONT_FAMILY
p2.font.color.rgb = POLAR_BLUE
p2.alignment = PP_ALIGN.CENTER

# ---------------------------------------------------------
# 2. LEFT: THE CHALLENGE
# ---------------------------------------------------------
# Outer Box
ch_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.95), Inches(3.4), Inches(4.55))
ch_box.fill.solid()
ch_box.fill.fore_color.rgb = WHITE
ch_box.line.color.rgb = BORDER_CARD
ch_box.line.width = Pt(1)

# Red Header
rh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.95), Inches(3.4), Inches(0.48))
rh.fill.solid()
rh.fill.fore_color.rgb = RED_HEADER
rh.line.fill.background()
tf = rh.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.text = "⚠  THE CHALLENGE"
p.font.size = Pt(12)
p.font.bold = True
p.font.name = FONT_FAMILY
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

# Challenge Items (4 items)
items = [
    ("📄 Data Silos", "Manifests & stock trapped in disconnected spreadsheets."),
    ("📦 Resupply Delays", "8-month winter isolation with zero resupply windows."),
    ("⛽ Cascading Failures", "Fuel shortage → Generator trip → Station heat loss."),
    ("📡 Total Blackout", "Polar blizzards cut all cloud & satellite links.")
]

for idx, (title, desc) in enumerate(items):
    iy = 1.55 + idx * 0.70
    ibox = slide.shapes.add_textbox(Inches(0.75), Inches(iy), Inches(3.1), Inches(0.65))
    tf = ibox.text_frame
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(10)
    p1.font.bold = True
    p1.font.name = FONT_FAMILY
    p1.font.color.rgb = NAVY_TEXT
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(8.5)
    p2.font.name = FONT_FAMILY
    p2.font.color.rgb = SLATE_TEXT

# Callout Banner at bottom of Challenge
callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(4.60), Inches(3.16), Inches(0.78))
callout.fill.solid()
callout.fill.fore_color.rgb = ORANGE_CALLOUT
callout.line.color.rgb = ORANGE_BORDER
callout.line.width = Pt(1)
tf = callout.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = tf.margin_right = Inches(0.10)
p = tf.paragraphs[0]
p.text = "\"Reactive tools report shortages.\nExtreme missions need cascading foresight.\""
p.font.size = Pt(8.5)
p.font.bold = True
p.font.name = FONT_FAMILY
p.font.color.rgb = RGBColor(180, 83, 9)
p.alignment = PP_ALIGN.CENTER

# ---------------------------------------------------------
# 3. CENTER: PHONE MOCKUP + 6 PROTOTYPE FEATURE BOXES
# ---------------------------------------------------------
# Center Phone Graphic
phone_path = os.path.abspath("assets/presentation/extracted/p2_img9_117_651x991.jpeg")
if os.path.exists(phone_path):
    slide.shapes.add_picture(phone_path, Inches(5.85), Inches(1.18), width=Inches(1.65), height=Inches(4.15))

# 6 Prototype Boxes (Precise 2-line layout)
boxes = [
    # (Left, Top, Width, Height, Title, Desc)
    (4.15, 1.10, 1.60, 0.95, "📊 Continuity Score", "0–100% live mission health index"),
    (7.60, 1.10, 1.60, 0.95, "🚢 Cargo Tracking", "Tracks Goa → Cape Town → Bases"),
    (4.15, 2.45, 1.60, 0.95, "🔄 What-If Simulator", "Tests delays & rationing in real-time"),
    (7.60, 2.45, 1.60, 0.95, "🧠 Mission Memory", "Recalls past expedition incident logs"),
    (4.15, 4.25, 1.60, 0.95, "🔗 Dependency Graph", "Links Fuel → Power → Heat → Labs"),
    (7.60, 4.25, 1.60, 0.95, "🛡️ Offline Triage", "Instant SOS dispatch & emergency logs"),
]

for (bx, by, bw, bh, btitle, bdesc) in boxes:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(by), Inches(bw), Inches(bh))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = BORDER_BLUE
    box.line.width = Pt(1.2)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.06)
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    p1 = tf.paragraphs[0]
    p1.text = btitle
    p1.font.size = Pt(9.2)
    p1.font.bold = True
    p1.font.name = FONT_FAMILY
    p1.font.color.rgb = POLAR_BLUE
    p1.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = bdesc
    p2.font.size = Pt(7.8)
    p2.font.name = FONT_FAMILY
    p2.font.color.rgb = SLATE_TEXT
    p2.alignment = PP_ALIGN.CENTER

# ---------------------------------------------------------
# 4. RIGHT: OUR SOLUTION
# ---------------------------------------------------------
# Outer Box
sol_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.35), Inches(0.95), Inches(3.4), Inches(4.55))
sol_box.fill.solid()
sol_box.fill.fore_color.rgb = WHITE
sol_box.line.color.rgb = BORDER_CARD
sol_box.line.width = Pt(1)

# Dark Blue Header
dh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.35), Inches(0.95), Inches(3.4), Inches(0.48))
dh.fill.solid()
dh.fill.fore_color.rgb = DARK_HEADER
dh.line.fill.background()
tf = dh.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.text = "🛡️  OUR SOLUTION"
p.font.size = Pt(12)
p.font.bold = True
p.font.name = FONT_FAMILY
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

# Sub-Header Banner
sub_b = slide.shapes.add_textbox(Inches(9.45), Inches(1.50), Inches(3.2), Inches(0.30))
tf = sub_b.text_frame
tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
p = tf.paragraphs[0]
p.text = "FROM REACTIVE TO PREDICTIVE"
p.font.size = Pt(9)
p.font.bold = True
p.font.name = FONT_FAMILY
p.font.color.rgb = RED_HEADER
p.alignment = PP_ALIGN.CENTER

# Column Headers
tbl_hdr = slide.shapes.add_textbox(Inches(9.45), Inches(1.78), Inches(3.2), Inches(0.25))
tf = tbl_hdr.text_frame
tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
p = tf.paragraphs[0]
r1 = p.add_run()
r1.text = "CONVENTIONAL         "
r1.font.bold = True
r1.font.underline = True
r1.font.size = Pt(8.5)
r1.font.name = FONT_FAMILY
r1.font.color.rgb = NAVY_TEXT

r2 = p.add_run()
r2.text = "POLAR-AI"
r2.font.bold = True
r2.font.underline = True
r2.font.size = Pt(8.5)
r2.font.name = FONT_FAMILY
r2.font.color.rgb = POLAR_BLUE

# Comparison Rows
rows = [
    ("Static spreadsheets", "0–100% Live health index"),
    ("Alerts after failure", "Predicts 7–14 days ahead"),
    ("Isolated cargo delays", "Cascading life-support trace"),
    ("Manual guesswork", "What-If scenario simulation"),
    ("Fails when offline", "100% Offline-first PWA")
]

for idx, (conv, pol) in enumerate(rows):
    ry = 2.10 + idx * 0.64
    rbox = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.45), Inches(ry), Inches(3.2), Inches(0.56))
    rbox.fill.solid()
    rbox.fill.fore_color.rgb = SKY_TINT if idx % 2 == 1 else WHITE
    rbox.line.color.rgb = BORDER_BLUE if idx % 2 == 1 else BORDER_CARD
    rbox.line.width = Pt(0.75)
    tf = rbox.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = f"{conv}  →  {pol}"
    p.font.size = Pt(7.8)
    p.font.name = FONT_FAMILY
    p.font.color.rgb = NAVY_TEXT
    p.alignment = PP_ALIGN.CENTER

# ---------------------------------------------------------
# 5. BOTTOM: REPLACED CORE USP CONTAINER
# ---------------------------------------------------------
# Main Outer Box replacing the lifecycle bar
usp_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.62), Inches(12.15), Inches(1.58))
usp_container.fill.solid()
usp_container.fill.fore_color.rgb = WHITE
usp_container.line.color.rgb = POLAR_BLUE
usp_container.line.width = Pt(1.5)

# Dark Navy Top Ribbon of USP
ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(5.62), Inches(12.15), Inches(0.32))
ribbon.fill.solid()
ribbon.fill.fore_color.rgb = DARK_HEADER
ribbon.line.fill.background()
tf = ribbon.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.text = "⭐  CORE USP: MISSION CONTINUITY INTELLIGENCE FOR EXTREME ENVIRONMENTS"
p.font.size = Pt(9.5)
p.font.bold = True
p.font.name = FONT_FAMILY
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

# Main USP Sentence
stmt_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.96), Inches(11.75), Inches(0.30))
tf = stmt_box.text_frame
tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
p = tf.paragraphs[0]
p.text = "“POLAR-AI connects cargo, stock, assets, and weather to predict cascading failures and simulate mitigation options before shortages become life-critical.”"
p.font.size = Pt(9.2)
p.font.bold = True
p.font.name = FONT_FAMILY
p.font.color.rgb = NAVY_TEXT
p.alignment = PP_ALIGN.CENTER

# 3 Horizontal Cards at bottom
usp_pillars = [
    ("1. CASCADING FAILURE TRACE", "Fuel Shortage → Generator Trip → Station Heat Loss", RED_HEADER),
    ("2. INTERACTIVE WHAT-IF ENGINE", "Tests resupply delays & rationing deltas before field action", POLAR_BLUE),
    ("3. 100% OFFLINE PWA & HUMAN GATE", "Zero-internet field autonomy; Commander retains final approval", GREEN_ACCENT),
]

for idx, (head, body, acc) in enumerate(usp_pillars):
    ux = 0.8 + idx * 3.98
    ucard = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ux), Inches(6.30), Inches(3.78), Inches(0.78))
    ucard.fill.solid()
    ucard.fill.fore_color.rgb = SKY_TINT
    ucard.line.color.rgb = BORDER_BLUE
    ucard.line.width = Pt(1)
    tf = ucard.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.10)
    p1 = tf.paragraphs[0]
    p1.text = head
    p1.font.size = Pt(8.5)
    p1.font.bold = True
    p1.font.name = FONT_FAMILY
    p1.font.color.rgb = acc
    p1.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(7.8)
    p2.font.name = FONT_FAMILY
    p2.font.color.rgb = NAVY_TEXT
    p2.alignment = PP_ALIGN.CENTER

# Save presentation
out_pptx = os.path.abspath("assets/presentation/revised_slide_2.pptx")
prs.save(out_pptx)
print("[SUCCESS] Created revised_slide_2.pptx")
